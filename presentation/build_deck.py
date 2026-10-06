"""Builds MACE_Agentic_Discovery.pptx with python-pptx. Run from the repo root:

    python presentation/build_deck.py
"""
import copy
import json
import math
import re
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_MARKER_STYLE
from pptx.enum.dml import MSO_LINE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

HERE = Path(__file__).parent
ROOT = HERE.parent
GA = json.loads((ROOT / "demo/backup/ga_history.json").read_text())

MAROON, MAROON2 = "500000", "7A2E2E"
NAVY, SLATE, MUTED = "1E293B", "64748B", "94A3B8"
ROSE, ROSE2, PANEL = "E3C4C4", "B87878", "F5F1F1"
GOLD, LINE, WHITE = "C9A227", "E2E8F0", "FFFFFF"
BODY, MATH = "Arial", "Cambria Math"

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]


def rgb(h):
    return RGBColor.from_string(h)


# ---------- text ----------
_TOK = re.compile(r"(_\{[^}]*\}|\^\{[^}]*\})")


def _runs(p, text, size, color, bold, italic, font):
    for part in _TOK.split(text):
        if not part:
            continue
        base = 0
        if part.startswith("_{"):
            part, base = part[2:-1], -25000
        elif part.startswith("^{"):
            part, base = part[2:-1], 30000
        r = p.add_run()
        r.text = part
        f = r.font
        f.size, f.bold, f.italic, f.name = Pt(size), bold, italic, font
        f.color.rgb = rgb(color)
        if base:
            r._r.get_or_add_rPr().set("baseline", str(base))


def text(slide, x, y, w, h, paras, size=14, color=NAVY, bold=False, italic=False,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, font=BODY, spacing=None, name=None):
    """paras: str or list of paragraphs; a paragraph is str or list of (text, overrides) runs."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    if isinstance(paras, str):
        paras = [paras]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if spacing:
            p.space_after = Pt(spacing)
        runs = [(para, {})] if isinstance(para, str) else para
        for t, o in runs:
            _runs(p, t, o.get("size", size), o.get("color", color), o.get("bold", bold),
                  o.get("italic", italic), o.get("font", font))
    return tb


# ---------- shapes ----------
def _no_shadow(sh):
    sp = sh._element.spPr
    for e in sp.findall(qn("a:effectLst")):
        sp.remove(e)
    sp.append(etree.SubElement(sp, qn("a:effectLst")))


def soft_shadow(sh):
    sp = sh._element.spPr
    for e in sp.findall(qn("a:effectLst")):
        sp.remove(e)
    eff = etree.SubElement(sp, qn("a:effectLst"))
    sh_ = etree.SubElement(eff, qn("a:outerShdw"), blurRad="114300", dist="25400", dir="5400000",
                           algn="t", rotWithShape="0")
    clr = etree.SubElement(sh_, qn("a:srgbClr"), val="1E293B")
    etree.SubElement(clr, qn("a:alpha"), val="14000")


def box(slide, x, y, w, h, fill=WHITE, line=None, radius=0.08, shape=None, shadow=False,
        line_w=1.0, dash=False, name=None):
    kind = shape or (MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE)
    sh = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        sh.name = name
    if kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        sh.adjustments[0] = min(0.5, radius / min(w, h))
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = rgb(fill)
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = rgb(line)
        sh.line.width = Pt(line_w)
        if dash:
            sh.line.dash_style = MSO_LINE.DASH
    soft_shadow(sh) if shadow else _no_shadow(sh)
    sh.text_frame.text = ""
    return sh


def label_in(sh, paras, size=14, color=NAVY, bold=False, align=PP_ALIGN.CENTER,
             anchor=MSO_ANCHOR.MIDDLE, font=BODY, italic=False, margin=0.08):
    tf = sh.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    tf.vertical_anchor = anchor
    if isinstance(paras, str):
        paras = [paras]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        runs = [(para, {})] if isinstance(para, str) else para
        for t, o in runs:
            _runs(p, t, o.get("size", size), o.get("color", color), o.get("bold", bold),
                  o.get("italic", italic), o.get("font", font))
    return sh


def circle(slide, cx, cy, d, fill, line=None, glyph=None, size=14, color=WHITE, font=BODY, line_w=1.5):
    sh = box(slide, cx - d / 2, cy - d / 2, d, d, fill=fill, line=line, shape=MSO_SHAPE.OVAL, line_w=line_w)
    if glyph:
        label_in(sh, glyph, size=size, color=color, bold=True, font=font, margin=0)
    return sh


def arrow(slide, x1, y1, x2, y2, color=MUTED, w=1.75, dash=False, head=True):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(w)
    if dash:
        c.line.dash_style = MSO_LINE.DASH
    if head:
        ln = c.line._get_or_add_ln()
        etree.SubElement(ln, qn("a:tailEnd"), type="triangle", w="med", len="med")
    return c


# ---------- slide frame ----------
N = [0]


def new_slide(kicker=None, title=None, subtitle=None, bg=WHITE, number=True):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = rgb(bg)
    N[0] += 1
    if kicker:
        text(s, 0.6, 0.38, 9, 0.3, kicker, size=11, bold=True, color=ROSE2, name="Kicker")
    if title:
        text(s, 0.6, 0.62, 12.1, 0.7, title, size=34, bold=True, color=MAROON, name="Title")
    if subtitle:
        text(s, 0.6, 1.32, 12.1, 0.4, subtitle, size=15, italic=True, color=SLATE, name="Subtitle")
    if number:
        text(s, 12.2, 7.0, 0.6, 0.25, f"{N[0]:02d}", size=10, bold=True,
             color=MAROON if bg == WHITE else ROSE, align=PP_ALIGN.RIGHT)
    return s


def pill(slide, msg, y=6.4, x=0.6, w=12.13):
    sh = box(slide, x, y, w, 0.5, fill=MAROON, radius=0.12)
    label_in(sh, msg, size=15, color=WHITE, bold=True)


def source(slide, msg):
    text(slide, 0.6, 7.03, 10.5, 0.25, msg, size=9, color=MUTED)


def notes(slide, msg):
    slide.notes_slide.notes_text_frame.text = msg


# =====================================================================
# 1. Title
# =====================================================================
s = new_slide(number=False)
pic = s.shapes.add_picture(str(HERE / "img/title_network.png"), Inches(6.3), Inches(0), height=Inches(7.5))
pic.crop_right = (pic.width - (prs.slide_width - pic.left)) / pic.width
pic.width = prs.slide_width - pic.left
blip = pic._element.find(".//" + qn("a:blip"))
etree.SubElement(blip, qn("a:alphaModFix"), amt="70000")
text(s, 0.8, 1.95, 7, 0.35, "AGENTIC DISCOVERY & DEVELOPMENT  ·  TEXAS A&M", size=14, bold=True, color=MAROON)
text(s, 0.75, 2.35, 7, 1.6, "MACE", size=110, bold=True, color=NAVY)
text(s, 0.8, 4.05, 5.3, 1.0, "Machine-learned atomic forces as an agent's physics engine",
     size=24, color=SLATE)
text(s, 0.8, 5.35, 6, 0.35, "Joshua Jang  ·  Fall 2026", size=14, italic=True, color=MUTED)
notes(s, "[0:00–0:15] Today: MACE, a machine-learned model of atomic forces, and why it is the kind of "
         "tool an AI agent needs to do real materials and chemistry discovery. Five minutes of slides, "
         "then a live demo where Claude Code drives MACE on this laptop.")

# =====================================================================
# 2. What is MACE + who is it for
# =====================================================================
s = new_slide("PART 1 · THE MODEL", "What is MACE?",
              "A machine-learned interatomic potential: quantum-level accuracy at force-field speed")
# input card with a little molecule
c_in = box(s, 0.6, 1.95, 3.3, 2.35, fill=PANEL, radius=0.15)
mol = [(1.55, 2.75, 0.42, MAROON), (2.3, 2.5, 0.32, ROSE2), (2.15, 3.2, 0.32, ROSE2), (1.0, 3.25, 0.32, NAVY),
       (2.95, 2.95, 0.28, MUTED)]
for a, b in [(0, 1), (0, 2), (0, 3), (1, 4), (2, 4)]:
    arrow(s, mol[a][0], mol[a][1], mol[b][0], mol[b][1], color=MUTED, w=2.5, head=False)
for x, y, d, col in mol:
    circle(s, x, y, d, col, line=WHITE)
text(s, 0.85, 3.6, 2.9, 0.3, "Atoms in", size=16, bold=True, color=NAVY)
text(s, 0.85, 3.92, 2.9, 0.3, "positions r_{i}  +  elements Z_{i}", size=12, color=SLATE)
arrow(s, 4.0, 3.12, 4.55, 3.12, color=MAROON, w=2.5)
m = box(s, 4.65, 1.95, 3.4, 2.35, fill=MAROON, radius=0.15, shadow=True)
label_in(m, [[("MACE", {"size": 36, "bold": True})],
             [("equivariant graph", {"size": 13})], [("neural network", {"size": 13})]], color=WHITE)
arrow(s, 8.15, 3.12, 8.7, 3.12, color=MAROON, w=2.5)
outs = [("E", "total energy", "eV"), ("F_{i}", "force on every atom", "eV/Å"), ("σ", "stress on the cell", "GPa")]
for k, (sym, what, unit) in enumerate(outs):
    y = 1.95 + k * 0.82
    b = box(s, 8.8, y, 3.95, 0.7, fill=WHITE, line=LINE, radius=0.12, shadow=True)
    circle(s, 9.2, y + 0.35, 0.5, MAROON if k == 1 else ROSE, glyph=sym, size=15,
           color=WHITE if k == 1 else MAROON, font=MATH)
    text(s, 9.62, y + 0.12, 3.0, 0.5, [[(what, {"bold": True, "color": NAVY}), (f"   {unit}", {"color": MUTED})]],
         size=14, anchor=MSO_ANCHOR.MIDDLE)
text(s, 0.6, 4.62, 6, 0.3, "WHO USES IT", size=11, bold=True, color=ROSE2)
who = [("Li", "Materials scientists", "batteries, catalysts, alloys"),
       ("C", "Chemists & pharma", "molecules, liquids, drug-like systems"),
       ("MD", "Simulation engineers", "10⁴–10⁶ atoms, nanoseconds of motion"),
       ("AI", "AI agents", "a fast, callable physics oracle")]
for k, (g, t, d) in enumerate(who):
    x = 0.6 + k * 3.1
    hl = k == 3
    box(s, x, 4.95, 2.85, 1.85, fill=MAROON if hl else PANEL, radius=0.15, shadow=hl)
    circle(s, x + 0.55, 5.45, 0.62, WHITE if hl else MAROON, glyph=g, size=15, color=MAROON if hl else WHITE)
    text(s, x + 0.25, 5.9, 2.45, 0.3, t, size=15, bold=True, color=WHITE if hl else NAVY)
    text(s, x + 0.25, 6.2, 2.45, 0.5, d, size=12, color=ROSE if hl else SLATE)
source(s, "Batatia et al., NeurIPS 2022 (MACE); MACE-MP-0 foundation model covers 89 elements")
notes(s, "[0:15–0:45] MACE is a machine-learned interatomic potential. You give it atoms (where they are and "
         "what element they are) and it returns the energy, the force on every atom, and the stress on the cell. "
         "It is trained on quantum-mechanical DFT calculations, so it is close to DFT accuracy but thousands of times "
         "faster. Users: materials scientists, chemists, simulation engineers, and increasingly AI agents that need "
         "a physics tool they can call.")

# =====================================================================
# 3. Accuracy vs cost
# =====================================================================
s = new_slide("PART 1 · THE MODEL", "Where MACE sits: accuracy vs. cost",
              "Learn from expensive quantum calculations once, then predict cheaply forever")
X0, Y0, X1, Y1 = 1.1, 1.95, 6.5, 5.9
arrow(s, X0, Y1, X1, Y1, color=SLATE, w=1.5)
arrow(s, X0, Y1, X0, Y0, color=SLATE, w=1.5)
text(s, X0, Y1 + 0.1, X1 - X0, 0.3, "compute cost per calculation  (log scale)  →", size=12, color=SLATE)
t = text(s, -0.55, Y0 + 1.75, 2.2, 0.35, "accuracy  →", size=12, color=SLATE, align=PP_ALIGN.CENTER,
         anchor=MSO_ANCHOR.MIDDLE)
t.rotation = -90
pts = [("Classical\nforce fields", 1.75, 5.15, 0.55, MUTED), ("MACE", 2.95, 2.95, 0.95, MAROON),
       ("DFT", 4.95, 2.75, 0.75, NAVY), ("Coupled\ncluster", 6.0, 2.2, 0.6, SLATE)]
for lab, x, y, d, col in pts:
    circle(s, x, y, d, col, line=WHITE)
arrow(s, 4.55, 2.82, 3.5, 2.92, color=GOLD, w=2.25, dash=True)
text(s, 3.45, 2.2, 1.6, 0.3, "trained on DFT", size=11, italic=True, color=GOLD, bold=True)
text(s, 1.15, 4.4, 1.5, 0.5, "Classical force fields", size=12, color=SLATE, align=PP_ALIGN.CENTER)
text(s, 2.25, 3.5, 1.4, 0.3, "MACE", size=14, bold=True, color=MAROON, align=PP_ALIGN.CENTER)
text(s, 4.35, 3.2, 1.2, 0.3, "DFT", size=13, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
text(s, 5.3, 2.6, 1.4, 0.5, "Coupled cluster", size=12, color=SLATE, align=PP_ALIGN.CENTER)
box(s, 3.3, 4.0, 3.0, 1.0, fill=PANEL, radius=0.12)
text(s, 3.45, 4.05, 2.75, 0.9, [[("Sweet spot: ", {"bold": True, "color": MAROON}),
                                  ("near-DFT accuracy at a tiny fraction of the cost", {})]],
     size=13, color=NAVY, anchor=MSO_ANCHOR.MIDDLE)
# scaling chart
cd = CategoryChartData()
cd.categories = ["10", "100", "1,000", "10,000"]
cd.add_series("DFT  ∝ N³", (1e3, 1e6, 1e9, 1e12))
cd.add_series("MACE  ∝ N", (1, 10, 100, 1000))
gf = s.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, Inches(7.1), Inches(1.85), Inches(5.65), Inches(4.35), cd)
ch = gf.chart


def style_chart(ch, title, ytitle=None, xtitle=None):
    ch.font.size, ch.font.name = Pt(11), BODY
    ch.font.color.rgb = rgb(SLATE)
    ch.has_title = True
    ch.chart_title.text_frame.text = title
    tr = ch.chart_title.text_frame.paragraphs[0].runs[0].font
    tr.size, tr.bold, tr.color.rgb = Pt(14), True, rgb(MAROON)
    ch.has_legend = True
    ch.legend.position = XL_LEGEND_POSITION.BOTTOM
    ch.legend.include_in_layout = False
    ch.legend.font.size = Pt(11)
    va, ca = ch.value_axis, ch.category_axis
    va.major_gridlines.format.line.color.rgb = rgb(LINE)
    va.format.line.fill.background()
    ca.format.line.color.rgb = rgb(MUTED)
    va.tick_labels.font.size = ca.tick_labels.font.size = Pt(11)
    for ax, t in ((va, ytitle), (ca, xtitle)):
        if t:
            ax.has_title = True
            ax.axis_title.text_frame.text = t
            f = ax.axis_title.text_frame.paragraphs[0].runs[0].font
            f.size, f.bold, f.color.rgb = Pt(11), False, rgb(SLATE)


def style_series(ser, color, width=2.75, dash=False, marker=True, msize=7):
    ser.format.line.color.rgb = rgb(color)
    ser.format.line.width = Pt(width)
    ser.smooth = False
    if dash:
        ser.format.line.dash_style = MSO_LINE.DASH
    if marker:
        ser.marker.style = XL_MARKER_STYLE.CIRCLE
        ser.marker.size = msize
        ser.marker.format.fill.solid()
        ser.marker.format.fill.fore_color.rgb = rgb(color)
        ser.marker.format.line.color.rgb = rgb(WHITE)
    else:
        ser.marker.style = XL_MARKER_STYLE.NONE


style_chart(ch, "Relative compute vs. number of atoms", "relative cost (log)", "atoms in the simulation")
style_series(ch.series[0], NAVY)
style_series(ch.series[1], MAROON)
scaling = ch.value_axis._element.find(qn("c:scaling"))
scaling.insert(0, etree.Element(qn("c:logBase"), val="10"))
ch.value_axis.tick_labels.number_format = "0E+0"
ch.value_axis.tick_labels.number_format_is_linked = False
box(s, 10.55, 3.25, 2.0, 0.6, fill=WHITE, line=ROSE, radius=0.12)
text(s, 10.6, 3.3, 1.9, 0.5, [[("10⁹× gap", {"bold": True, "color": MAROON, "size": 15})],
                              [("at 10k atoms", {"size": 10, "color": SLATE})]],
     align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
pill(s, "MACE scales linearly: big systems and long simulations become possible")
source(s, "Scaling curves are illustrative (DFT ~ O(N³), local ML potentials ~ O(N)); constants vary by code and hardware")
notes(s, "[0:45–1:15] Trade-off map. Classical force fields are fast but crude. DFT is accurate but expensive and "
         "scales as N cubed. Coupled cluster is the gold standard and even slower. MACE learns from DFT once, then "
         "predicts at a tiny fraction of the cost, and it scales linearly with the number of atoms, so the gap grows "
         "with system size.")

# =====================================================================
# 4. Physics: forces are the slope
# =====================================================================
s = new_slide("PART 1 · THE MODEL", "Forces are the slope of the energy",
              "Learn one function, the energy, and every force follows by differentiation")
s.shapes.add_picture(str(HERE / "img/pes.png"), Inches(0.5), Inches(1.85), width=Inches(6.1))
# card 1: locality with mini cutoff diagram
box(s, 6.95, 1.95, 5.8, 2.05, fill=PANEL, radius=0.15)
cx, cy = 7.95, 2.97
box(s, cx - 0.75, cy - 0.75, 1.5, 1.5, fill=None, line=MAROON, shape=MSO_SHAPE.OVAL, dash=True, line_w=1.25)
for ang, r_ in [(20, 0.5), (110, 0.55), (200, 0.45), (290, 0.55), (155, 1.0), (235, 1.02)]:
    px, py = cx + r_ * math.cos(math.radians(ang)), cy + r_ * math.sin(math.radians(ang))
    circle(s, px, py, 0.22, ROSE2 if r_ < 0.75 else LINE, line=WHITE)
circle(s, cx, cy, 0.32, MAROON, line=WHITE)
text(s, cx + 0.12, cy - 0.98, 0.8, 0.25, "r_{c} ≈ 6 Å", size=10, color=MAROON, bold=True)
text(s, 9.05, 2.2, 3.6, 0.55, "E = Σ_{i} E_{i}( neighbors within r_{c} )", size=19, font=MATH, color=NAVY)
text(s, 9.05, 2.95, 3.55, 0.9, "Total energy is a sum of local atomic contributions, so cost grows linearly",
     size=13, color=SLATE)
# card 2: forces
box(s, 6.95, 4.2, 5.8, 2.05, fill=PANEL, radius=0.15)
circle(s, 7.95, 5.22, 1.1, MAROON, glyph="−∇E", size=22, font=MATH)
text(s, 9.05, 4.45, 3.6, 0.55, "F_{i} = − ∂E / ∂r_{i}", size=22, font=MATH, color=NAVY)
text(s, 9.05, 5.2, 3.55, 0.9, "Forces by automatic differentiation: energy is conserved, so dynamics stay stable",
     size=13, color=SLATE)
pill(s, "Train on DFT energies + forces  →  predict both for brand-new structures in milliseconds")
notes(s, "[1:15–1:55] The core physics. Atoms sit on a potential energy surface. The force is the negative slope: "
         "stretch a bond and the force pulls it back; at the minimum the force is zero. MACE writes the total energy as "
         "a sum of atomic energies, each depending only on neighbors within a cutoff of about 6 angstroms. Forces come "
         "from differentiating that energy automatically, which guarantees energy conservation in molecular dynamics.")

# =====================================================================
# 5. Architecture / math
# =====================================================================
s = new_slide("PART 1 · THE MODEL", "Inside MACE: many-body equivariant messages",
              "A graph neural network where atoms are nodes and neighbors within r_{c} are edges")
steps = [("Embed", "h_{j} = W · onehot(Z_{j})", "each element becomes a learned vector"),
         ("Neighbor basis", "A_{i} = Σ_{j} R(r_{ij}) Y_{lm}(r̂_{ij}) h_{j}", "distance × direction (spherical harmonics)"),
         ("Many-body product", "B_{i} = A_{i} ⊗ A_{i} ⊗ A_{i}", "tensor products build 4-body terms in one step"),
         ("Update + readout", "E_{i} = MLP(h_{i}^{(2)})", "two layers, then a per-atom energy")]
for k, (t, eq, d) in enumerate(steps):
    x = 0.6 + k * 3.1
    box(s, x, 1.95, 2.85, 2.95, fill=PANEL, radius=0.15)
    circle(s, x + 1.425, 2.55, 0.72, MAROON, glyph=str(k + 1), size=22)
    text(s, x + 0.15, 3.1, 2.55, 0.35, t, size=16, bold=True, color=MAROON, align=PP_ALIGN.CENTER)
    text(s, x + 0.12, 3.55, 2.61, 0.5, eq, size=14, font=MATH, color=NAVY, align=PP_ALIGN.CENTER)
    text(s, x + 0.2, 4.1, 2.45, 0.65, d, size=12, color=SLATE, align=PP_ALIGN.CENTER)
    if k < 3:
        arrow(s, x + 2.88, 3.42, x + 3.07, 3.42, color=MAROON, w=2.25)
chips = [("E(3)", "Equivariant", "rotate the molecule → forces rotate with it"),
         ("ν=3", "High body order", "rich physics with only 2 layers"),
         ("2L", "Small + fast", "local, parallel, GPU-friendly")]
for k, (g, t, d) in enumerate(chips):
    x = 0.6 + k * 4.12
    box(s, x, 5.2, 3.9, 1.15, fill=WHITE, line=LINE, radius=0.15, shadow=True)
    circle(s, x + 0.6, 5.775, 0.78, ROSE, glyph=g, size=13, color=MAROON, font=MATH)
    text(s, x + 1.15, 5.33, 2.65, 0.35, t, size=15, bold=True, color=NAVY)
    text(s, x + 1.15, 5.68, 2.65, 0.6, d, size=12, color=SLATE)
source(s, "Batatia, Kovács, Simm, Ortner, Csányi, “MACE: Higher Order Equivariant Message Passing Neural Networks,” NeurIPS 2022")
notes(s, "[1:55–2:40] The engineering. MACE is a graph neural network: atoms are nodes, neighbors within the cutoff are "
         "edges. Step 1 embeds each element. Step 2 builds a neighbor basis from distances (radial functions) and "
         "directions (spherical harmonics). Step 3 is the key idea from the Atomic Cluster Expansion: tensor products "
         "of that basis create 4-body interactions in a single message. Step 4 updates and reads out per-atom energy. "
         "Everything is equivariant, so rotating the input rotates the forces correctly, and high body order means only "
         "two layers are needed, which makes it fast.")

# =====================================================================
# 6. Agent loop + funnel
# =====================================================================
s = new_slide("PART 2 · THE AGENT", "MACE as the agent's physics tool",
              "The LLM plans and reasons; MACE answers “what would the atoms do?” in milliseconds")
hub = (3.75, 4.15)
circle(s, *hub, 1.45, WHITE, line=MAROON, glyph=[[("LLM", {"size": 18, "bold": True})], [("agent", {"size": 12})]],
       color=MAROON, line_w=2)
nodes = [("Plan", "hypothesis", 3.75, 2.3, False), ("Build", "ASE · pymatgen", 5.95, 4.15, False),
         ("Simulate", "MACE: relax · MD", 3.75, 6.0, True), ("Analyze", "rank · decide", 1.6, 4.15, False)]
for t, d, x, y, hl in nodes:
    b = box(s, x - 1.0, y - 0.4, 2.0, 0.8, fill=MAROON if hl else PANEL, radius=0.15, shadow=hl)
    label_in(b, [[(t, {"bold": True, "size": 14})], [(d, {"size": 10})]], color=WHITE if hl else NAVY)
for (x1, y1, x2, y2) in [(4.8, 2.5, 5.75, 3.7), (5.75, 4.6, 4.8, 5.8), (2.7, 5.8, 1.8, 4.6), (1.8, 3.7, 2.7, 2.5)]:
    arrow(s, x1, y1, x2, y2, color=MAROON, w=2)
fun = [("10⁶", "ideas generated", "agent / generative model", ROSE, NAVY, 5.2),
       ("10⁴", "screened with MACE", "milliseconds each", ROSE2, WHITE, 4.3),
       ("10²", "checked with DFT", "hours each", MAROON2, WHITE, 3.4),
       ("10", "made in the lab", "weeks each", MAROON, WHITE, 2.5)]
for k, (n, t, d, f, tc, w) in enumerate(fun):
    y = 1.95 + k * 1.05
    b = box(s, 10.15 - w / 2, y, w, 0.9, fill=f, radius=0.12)
    label_in(b, [[(n + "  ", {"size": 20, "bold": True}), (t, {"size": 13, "bold": True})],
                 [(d, {"size": 10})]], color=tc)
pill(s, "Physics cheap enough for an agent to think with", x=7.55, w=5.2)
notes(s, "[2:40–3:20] Where agents come in. The LLM is good at planning and reasoning but cannot know what atoms will "
         "do. MACE is a tool it can call: build a structure, relax it or run dynamics, read the energies and forces, decide "
         "what to try next. In a discovery funnel, the agent can generate a million ideas, MACE screens ten thousand in "
         "milliseconds each, DFT checks the top hundred, and only ten go to the lab.")

# =====================================================================
# 7. Evolution
# =====================================================================
s = new_slide("PART 2 · THE AGENT", "Evolution with MACE as the fitness function",
              "Generate, score, select, repeat: the LLM can be the mutation operator")
ring = [("Population", "candidate structures", 0), ("Mutate", "GA swaps  or  LLM ideas", 1),
        ("Score", "MACE energy", 2), ("Select", "keep the fittest", 3)]
cx, cy, R = 3.6, 3.75, 1.45
pos = []
for t, d, k in ring:
    ang = math.radians(-90 + k * 90)
    pos.append((cx + 1.85 * math.cos(ang), cy + R * math.sin(ang)))
for k, ((t, d, _), (x, y)) in enumerate(zip(ring, pos)):
    hl = t == "Score"
    b = box(s, x - 1.05, y - 0.4, 2.1, 0.8, fill=MAROON if hl else PANEL, radius=0.15, shadow=hl)
    label_in(b, [[(t, {"bold": True, "size": 14})], [(d, {"size": 10})]], color=WHITE if hl else NAVY)
for a, b_ in [(0, 1), (1, 2), (2, 3), (3, 0)]:
    (x1, y1), (x2, y2) = pos[a], pos[b_]
    dx, dy = x2 - x1, y2 - y1
    arrow(s, x1 + dx * 0.36, y1 + dy * 0.36, x1 + dx * 0.64, y1 + dy * 0.64, color=MAROON, w=2)
text(s, cx - 0.8, cy - 0.2, 1.6, 0.4, "×  generations", size=12, italic=True, color=SLATE, align=PP_ALIGN.CENTER)
# genome strip: L1_0 CuAu
text(s, 0.6, 5.8, 6, 0.3, "GENOME: 32 SITES, 4 LAYERS  (shown: L1₀ CuAu, the known stable phase)",
     size=10, bold=True, color=ROSE2)
for k, ch_ in enumerate("CCCCCCCC" "AAAAAAAA" "CCCCCCCC" "AAAAAAAA"):
    x = 0.6 + k * 0.18 + (k // 8) * 0.12
    box(s, x, 6.12, 0.15, 0.38, fill="B87333" if ch_ == "C" else GOLD, radius=0.03)
text(s, 0.6, 6.6, 6.5, 0.3, [[("■ ", {"color": "B87333"}), ("Cu    ", {}), ("■ ", {"color": GOLD}), ("Au", {})]],
     size=11, color=SLATE)
# chart from the real GA run
hist = GA["history"]
cd = CategoryChartData()
cd.categories = [str(h["gen"]) for h in hist]
cd.add_series("population mean", [round(h["mean"], 1) for h in hist])
cd.add_series("GA best", [round(h["best"], 1) for h in hist])
cd.add_series("L1₀ (agent's proposal)", [round(GA["L10"], 1)] * len(hist))
gf = s.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, Inches(7.0), Inches(1.9), Inches(5.75), Inches(4.75), cd)
ch = gf.chart
style_chart(ch, "Blind GA: +18   ·   agent's L1₀: −18 meV/atom", "formation energy (meV/atom)", "generation")
style_series(ch.series[0], ROSE, width=2, msize=5)
style_series(ch.series[1], MAROON)
style_series(ch.series[2], NAVY, width=2.25, dash=True, marker=False)
ch.value_axis.minimum_scale, ch.value_axis.maximum_scale = -30, 90
ch.value_axis.major_unit = 30
from pptx.enum.chart import XL_TICK_LABEL_POSITION
ch.category_axis.tick_label_position = XL_TICK_LABEL_POSITION.LOW
source(s, "Real run from today's demo: 32-atom Cu/Au cell, MACE-MP-0 (small), 235 MACE calls in 17 s on a laptop CPU")
notes(s, "[3:20–4:00] An evolutionary workflow. A population of candidate structures, mutate them, score every child "
         "with MACE, keep the fittest, repeat. The chart is a real run on this laptop: 32-atom copper-gold alloys, "
         "235 MACE calls in 17 seconds. The blind GA improves a lot, from about +50 to +18 meV per atom, but it stalls. "
         "The known stable phase, L1-zero, is at −18. That gap is the opening for an LLM: replace random mutation with "
         "an agent that proposes chemically sensible changes. You'll see exactly this in the demo.")

# =====================================================================
# 8. More patterns
# =====================================================================
s = new_slide("PART 2 · THE AGENT", "More ways to pair MACE with an agent",
              "Same tool, four workflow shapes")
pats = [("Active learning", ["MACE unsure?", "DFT labels it", "fine-tune MACE"], "the agent decides where the model needs more data"),
        ("MCP tool server", ["Claude", "relax · md · phonons", "MACE"], "expose MACE as tools any agent can call"),
        ("Multi-agent team", ["Planner", "Simulator (MACE)", "Critic"], "one proposes, one computes, one checks"),
        ("Property prediction", ["Agent sets T, P", "MACE dynamics", "diffusion, strength"], "long simulations that DFT can't afford")]
for k, (t, flow, d) in enumerate(pats):
    x = 0.6 + (k % 2) * 6.18
    y = 1.95 + (k // 2) * 2.45
    box(s, x, y, 5.95, 2.2, fill=PANEL, radius=0.15)
    circle(s, x + 0.5, y + 0.5, 0.52, MAROON, glyph=str(k + 1), size=15)
    text(s, x + 0.95, y + 0.33, 4.8, 0.35, t, size=17, bold=True, color=NAVY)
    for j, f in enumerate(flow):
        fx = x + 0.3 + j * 1.88
        hl = "MACE" in f
        b = box(s, fx, y + 0.98, 1.6, 0.6, fill=MAROON if hl else WHITE, line=None if hl else ROSE, radius=0.1)
        label_in(b, f, size=11, color=WHITE if hl else NAVY, bold=hl, margin=0.04)
        if j < 2:
            arrow(s, fx + 1.62, y + 1.28, fx + 1.86, y + 1.28, color=MAROON, w=1.75)
    text(s, x + 0.3, y + 1.72, 5.4, 0.35, d, size=12, italic=True, color=SLATE)
notes(s, "[4:00–4:30] Other patterns. Active learning: the agent watches MACE's uncertainty and sends only the uncertain "
         "structures to DFT, then fine-tunes. An MCP server wraps MACE as tools that Claude or any agent can call. A "
         "multi-agent team splits planning, simulation, and critique. And for properties like diffusion or mechanical "
         "strength, the agent sets up long MACE molecular-dynamics runs that DFT could never afford.")

# =====================================================================
# 9. Pitfalls
# =====================================================================
s = new_slide("PART 2 · THE AGENT", "Trust, but verify",
              "Pressure-test MACE before an agent acts on it")
pit = [("?", "Out of distribution", "Chemistry far from the training data fails silently. Watch uncertainty."),
       ("DFT", "Inherits DFT errors", "Trained on PBE-level DFT, so its known biases carry over."),
       ("≈", "Softened surface", "Universal models tend to under-predict stiffness. Fine-tune for production."),
       ("r_{c}", "Local by design", "A ~6 Å cutoff misses long-range electrostatics without add-ons.")]
for k, (g, t, d) in enumerate(pit):
    x = 0.6 + k * 3.1
    box(s, x, 1.95, 2.85, 3.5, fill=WHITE, line=LINE, radius=0.15, shadow=True)
    circle(s, x + 0.75, 2.75, 0.95, MAROON, glyph=g, size=18 if len(g) < 3 else 15, font=MATH)
    text(s, x + 0.3, 3.55, 2.35, 0.4, t, size=17, bold=True, color=NAVY)
    text(s, x + 0.3, 4.05, 2.35, 1.3, d, size=14, color=SLATE)
pill(s, "Agent rule: MACE proposes, DFT disposes. Spot-check every winner.", y=5.85)
notes(s, "[4:30–4:50] Four cautions before an agent trusts MACE. It can fail silently outside its training data. It "
         "inherits the errors of the DFT it learned from. Universal models are known to be a bit too soft. And it is local, "
         "so long-range electrostatics need extra terms. The rule for an agent: MACE screens, DFT confirms the winners.")

# =====================================================================
# 10. Demo
# =====================================================================
s = new_slide(bg=MAROON)
text(s, 0.8, 1.0, 11, 0.35, "LIVE DEMO", size=14, bold=True, color=ROSE)
text(s, 0.8, 1.4, 11.5, 1.1, "Claude Code  ×  MACE", size=54, bold=True, color=WHITE)
text(s, 0.8, 2.55, 11.5, 0.5, "An LLM agent in VS Code using MACE as its physics tool, live on this laptop",
     size=18, color=ROSE)
acts = [("1", "Forces", "Stretch a water molecule, read the forces, watch it relax"),
        ("2", "Blind evolution", "A genetic algorithm scores 235 Cu/Au alloys in 17 s"),
        ("3", "Agent evolution", "Claude proposes structures, MACE judges them")]
for k, (n, t, d) in enumerate(acts):
    x = 0.8 + k * 4.0
    box(s, x, 3.6, 3.7, 2.6, fill=WHITE, radius=0.15, shadow=True)
    circle(s, x + 0.65, 4.25, 0.75, MAROON, glyph=n, size=22)
    text(s, x + 0.3, 4.85, 3.1, 0.4, t, size=19, bold=True, color=MAROON)
    text(s, x + 0.3, 5.3, 3.1, 0.8, d, size=13, color=SLATE)
notes(s, "[4:50–5:00] Switch to VS Code. Follow DEMO_SCRIPT.md: Act 1 forces, Act 2 blind GA, Act 3 Claude as the "
         "mutation operator.")

# =====================================================================
# 11. Takeaways
# =====================================================================
s = new_slide(None, "Key takeaways", "Four ideas to leave with")
tk = [("E→F", "One model, all forces", "Learn the energy; forces and stress come free by differentiation."),
      ("ν", "Physics built in", "Equivariant, many-body messages make it accurate and fast."),
      ("AI", "Agents need cheap physics", "MACE turns an LLM's hypotheses into numbers in milliseconds."),
      ("✓", "Close the loop", "Evolve with MACE, verify winners with DFT and the lab.")]
for k, (g, t, d) in enumerate(tk):
    x = 0.6 + k * 3.1
    box(s, x, 1.95, 2.85, 3.6, fill=WHITE, line=LINE, radius=0.15, shadow=True)
    circle(s, x + 0.75, 2.75, 0.95, PANEL, glyph=g, size=16, color=MAROON, font=MATH)
    text(s, x + 2.0, 2.3, 0.6, 0.7, str(k + 1), size=40, bold=True, color=ROSE, align=PP_ALIGN.RIGHT)
    text(s, x + 0.3, 3.5, 2.35, 0.75, t, size=17, bold=True, color=NAVY, anchor=MSO_ANCHOR.BOTTOM)
    text(s, x + 0.3, 4.4, 2.35, 1.1, d, size=13, color=SLATE)
text(s, 0.6, 5.95, 12, 0.8, "Questions?", size=36, bold=True, color=MAROON)
notes(s, "[after demo, ~20 s] Recap: MACE learns the energy and gets forces for free; its equivariant many-body design is "
         "why it's accurate and fast; agents need exactly this kind of cheap, callable physics; and the loop closes with "
         "DFT and lab verification. Questions?")

out = HERE / "MACE_Agentic_Discovery.pptx"
prs.save(out)
print("saved", out)
